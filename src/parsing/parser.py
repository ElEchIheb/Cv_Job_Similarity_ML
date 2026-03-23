from __future__ import annotations

import re
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Dict, List
from xml.etree import ElementTree as ET

try:
    import pdfplumber  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    pdfplumber = None

try:
    from PyPDF2 import PdfReader  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    PdfReader = None

try:
    from docx import Document  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    Document = None

try:
    import spacy  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    spacy = None


_STOPWORDS = {
    "a",
    "about",
    "above",
    "after",
    "again",
    "against",
    "ai",
    "alors",
    "an",
    "and",
    "any",
    "are",
    "as",
    "at",
    "au",
    "aucun",
    "aussi",
    "autre",
    "aux",
    "avec",
    "avoir",
    "be",
    "been",
    "before",
    "below",
    "between",
    "both",
    "but",
    "by",
    "ce",
    "cela",
    "ces",
    "cet",
    "cette",
    "comme",
    "comment",
    "dans",
    "de",
    "dedans",
    "dehors",
    "des",
    "did",
    "do",
    "does",
    "doing",
    "don",
    "down",
    "during",
    "elle",
    "en",
    "entre",
    "est",
    "et",
    "etc",
    "eu",
    "eux",
    "for",
    "from",
    "further",
    "had",
    "has",
    "have",
    "he",
    "her",
    "here",
    "hers",
    "herself",
    "him",
    "himself",
    "his",
    "how",
    "i",
    "if",
    "il",
    "ils",
    "in",
    "into",
    "is",
    "it",
    "its",
    "itself",
    "je",
    "just",
    "la",
    "le",
    "les",
    "leur",
    "leurs",
    "lui",
    "ma",
    "mais",
    "me",
    "mes",
    "moi",
    "mon",
    "more",
    "most",
    "my",
    "myself",
    "ne",
    "no",
    "nor",
    "not",
    "nos",
    "notre",
    "nous",
    "of",
    "off",
    "on",
    "once",
    "only",
    "or",
    "other",
    "our",
    "ours",
    "ourselves",
    "out",
    "over",
    "own",
    "par",
    "pas",
    "pour",
    "plus",
    "qu",
    "que",
    "qui",
    "sa",
    "same",
    "se",
    "ses",
    "she",
    "should",
    "so",
    "some",
    "son",
    "sont",
    "sous",
    "sur",
    "t",
    "ta",
    "te",
    "than",
    "that",
    "the",
    "their",
    "theirs",
    "them",
    "themselves",
    "then",
    "there",
    "these",
    "they",
    "this",
    "those",
    "through",
    "to",
    "toi",
    "ton",
    "too",
    "tout",
    "tu",
    "under",
    "until",
    "up",
    "very",
    "via",
    "votre",
    "vous",
    "was",
    "we",
    "were",
    "what",
    "when",
    "where",
    "which",
    "while",
    "who",
    "whom",
    "why",
    "with",
    "you",
    "your",
    "yours",
    "yourself",
    "yourselves",
}

_SECTION_PATTERNS = {
    "summary": re.compile(
        r"^(summary|profil|profile|about|objectif|objective|resume|presentation)\s*:?\s*$",
        re.IGNORECASE,
    ),
    "experience": re.compile(
        r"^(experience|work experience|professional experience|career|employment|experiences?|"
        r"experience professionnelle|parcours professionnel|historique professionnel)\s*:?\s*$",
        re.IGNORECASE,
    ),
    "education": re.compile(
        r"^(education|formation|academic background|studies|diplomes?|certifications? acad[ée]miques?)\s*:?\s*$",
        re.IGNORECASE,
    ),
    "skills": re.compile(
        r"^(skills|competences?|technical skills|stack|outils|expertise|technologies)\s*:?\s*$",
        re.IGNORECASE,
    ),
    "languages": re.compile(
        r"^(languages?|langues?)\s*:?\s*$",
        re.IGNORECASE,
    ),
    "certifications": re.compile(
        r"^(certifications?|licenses?|licences?)\s*:?\s*$",
        re.IGNORECASE,
    ),
    "projects": re.compile(
        r"^(projects?|projets?)\s*:?\s*$",
        re.IGNORECASE,
    ),
}

_SPACY_MODEL_CACHE = None


def _strip_artifacts(text: str) -> str:
    cleaned = text.replace("\x00", " ")
    cleaned = re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "", cleaned)
    cleaned = cleaned.replace("•", "\n- ").replace("▪", "\n- ").replace("●", "\n- ")
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def _group_words_by_lines(words: List[dict]) -> List[str]:
    lines: Dict[int, List[dict]] = defaultdict(list)
    for word in words:
        line_key = int(round(float(word.get("top", 0.0)) / 3.0))
        lines[line_key].append(word)
    ordered_lines = []
    for line_key in sorted(lines):
        parts = [item.get("text", "").strip() for item in sorted(lines[line_key], key=lambda item: float(item.get("x0", 0.0)))]
        if parts:
            ordered_lines.append(" ".join(part for part in parts if part))
    return ordered_lines


def _extract_pdfplumber_page_text(page) -> str:  # pragma: no cover - depends on optional dependency
    words = page.extract_words(use_text_flow=True, keep_blank_chars=False)
    if not words:
        return page.extract_text() or ""

    width = float(getattr(page, "width", 0.0) or 0.0)
    left_column = [word for word in words if float(word.get("x0", 0.0)) < width * 0.55]
    right_column = [word for word in words if float(word.get("x0", 0.0)) >= width * 0.55]

    if width and left_column and right_column and len(right_column) >= max(3, len(words) * 0.15):
        left_lines = _group_words_by_lines(left_column)
        right_lines = _group_words_by_lines(right_column)
        return "\n".join(left_lines + right_lines)

    return "\n".join(_group_words_by_lines(words))


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF using pdfplumber first and PyPDF2 as a fallback."""

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PDF not found: {file_path}")

    text = ""
    if pdfplumber is not None:
        with pdfplumber.open(path) as pdf:
            pages = [_extract_pdfplumber_page_text(page) for page in pdf.pages]
        text = "\n\n".join(page for page in pages if page)
    elif PdfReader is not None:
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n\n".join(pages)
    else:  # pragma: no cover - fallback for environments without PDF libs
        raw = path.read_bytes()
        text_fragments = re.findall(rb"\(([^()]*)\)", raw)
        text = "\n".join(fragment.decode("latin1", errors="ignore") for fragment in text_fragments)

    return _strip_artifacts(text)


def _extract_docx_with_zipfile(file_path: str) -> str:
    with zipfile.ZipFile(file_path) as archive:
        xml_bytes = archive.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    texts = []
    for node in root.iter():
        if node.tag.endswith("}t") and node.text:
            texts.append(node.text)
        elif node.tag.endswith("}tr"):
            texts.append("\n")
        elif node.tag.endswith("}p"):
            texts.append("\n")
    return _strip_artifacts(" ".join(texts))


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX file, including table content."""

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {file_path}")

    if Document is not None:  # pragma: no branch - simple optional path
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return _strip_artifacts("\n".join(parts))

    return _extract_docx_with_zipfile(str(path))


def clean_text(text: str) -> str:
    """Remove contact artifacts while keeping punctuation useful for NLP."""

    cleaned = text or ""
    cleaned = cleaned.replace("\r", "\n")
    cleaned = re.sub(r"https?://\S+|www\.\S+", " ", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\b[\w.\-+%]+@[\w.\-]+\.[A-Za-z]{2,}\b", " ", cleaned)
    cleaned = re.sub(r"(?<!\w)(?:\+?\d[\d().\-\s]{7,}\d)", " ", cleaned)
    cleaned = re.sub(r"[ \t]+", " ", cleaned)
    cleaned = re.sub(r"\n\s+\n", "\n\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    cleaned = re.sub(r"[^\S\n]{2,}", " ", cleaned)
    return cleaned.strip().lower()


def _naive_lemma(token: str) -> str:
    if len(token) <= 4 or any(char in token for char in "+#./-"):
        return token
    for suffix in ("ments", "ment", "ation", "ations", "ingly", "ingly", "ingly", "ing", "edly", "edly", "ed", "ies", "es", "s"):
        if token.endswith(suffix) and len(token) - len(suffix) >= 3:
            if suffix == "ies":
                return token[:-3] + "y"
            return token[: -len(suffix)]
    return token


def _load_spacy_model():  # pragma: no cover - exercised only when spaCy is installed
    global _SPACY_MODEL_CACHE
    if _SPACY_MODEL_CACHE is not None:
        return _SPACY_MODEL_CACHE
    if spacy is None:
        return None
    for model_name in ("fr_core_news_lg", "en_core_web_lg", "fr_core_news_sm", "en_core_web_sm"):
        try:
            _SPACY_MODEL_CACHE = spacy.load(model_name)
            return _SPACY_MODEL_CACHE
        except Exception:
            continue
    return None


def tokenize_and_normalize(text: str) -> List[str]:
    """Tokenize text, lemmatize, remove stopwords and short tokens."""

    normalized_text = clean_text(text)
    nlp = _load_spacy_model()
    if nlp is not None:  # pragma: no cover - requires spaCy models
        doc = nlp(normalized_text)
        return [
            token.lemma_.lower().strip()
            for token in doc
            if token.is_alpha and not token.is_stop and len(token.lemma_.strip()) >= 2
        ]

    tokens = re.findall(r"[a-zA-Z][a-zA-Z0-9+#./-]{1,}", normalized_text)
    normalized_tokens = []
    for token in tokens:
        lemma = _naive_lemma(token)
        if lemma in _STOPWORDS or len(lemma) < 2:
            continue
        normalized_tokens.append(lemma)
    return normalized_tokens


def segment_cv(text: str) -> Dict[str, str]:
    """Split a CV into coarse sections using multilingual headings."""

    sections = {
        "summary": [],
        "experience": [],
        "education": [],
        "skills": [],
        "languages": [],
        "certifications": [],
        "projects": [],
        "other": [],
    }
    current_section = "summary"
    lines = [line.strip() for line in (text or "").splitlines()]

    for line in lines:
        if not line:
            continue
        matched_section = None
        for section_name, pattern in _SECTION_PATTERNS.items():
            if pattern.match(line):
                matched_section = section_name
                break
        if matched_section is not None:
            current_section = matched_section
            continue
        sections[current_section].append(line)

    if not sections["skills"]:
        skill_lines = [line for line in lines if any(token in line.lower() for token in ("python", "sql", "docker", "react", "comp", "skill"))]
        sections["skills"].extend(skill_lines)
    if not sections["experience"]:
        experience_lines = [line for line in lines if re.search(r"\b(20\d{2}|19\d{2}|ans|years?)\b", line, flags=re.IGNORECASE)]
        sections["experience"].extend(experience_lines[:20])
    if not sections["education"]:
        education_lines = [line for line in lines if re.search(r"\b(master|bachelor|licence|phd|doctorat|universit|ecole|school)\b", line, flags=re.IGNORECASE)]
        sections["education"].extend(education_lines[:10])
    if not sections["languages"]:
        language_lines = [line for line in lines if re.search(r"\b(english|french|francais|anglais|spanish|arabic)\b", line, flags=re.IGNORECASE)]
        sections["languages"].extend(language_lines[:10])

    return {name: "\n".join(values).strip() for name, values in sections.items()}

